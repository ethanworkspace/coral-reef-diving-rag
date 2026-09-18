from __future__ import annotations

import hashlib
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class ExtractedPart:
    label: str
    text: str


class _VisibleText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data.strip():
            self.parts.append(data.strip())


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract(path: Path) -> list[ExtractedPart]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        pages: list[ExtractedPart] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(ExtractedPart(f"page {index}", text))
        return pages
    if suffix == ".docx":
        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                text = " | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells)
                if text.strip(" |"):
                    paragraphs.append(text)
        return [ExtractedPart("document", "\n".join(paragraphs))]
    if suffix in {".txt", ".md", ".yaml", ".yml", ".json", ".csv"}:
        return [ExtractedPart("text", path.read_text(encoding="utf-8", errors="replace"))]
    if suffix in {".html", ".htm"}:
        parser = _VisibleText()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return [ExtractedPart("webpage snapshot", "\n".join(parser.parts))]
    return []


def chunk(parts: list[ExtractedPart], size: int = 900, overlap: int = 150) -> list[ExtractedPart]:
    output: list[ExtractedPart] = []
    for part in parts:
        text = " ".join(part.text.split())
        if not text:
            continue
        start = 0
        number = 1
        while start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                boundary = text.rfind("。", start, end)
                boundary = max(boundary, text.rfind(".", start, end), text.rfind(" ", start, end))
                if boundary > start + size // 2:
                    end = boundary + 1
            output.append(ExtractedPart(f"{part.label}, chunk {number}", text[start:end]))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
            number += 1
    return output
